"""
Тесты выгрузки таблицы и ответа сервера для расширения браузера.

Проверяется формат JSON из docs/json_format.md расширения, файл Excel и
разбор таблицы, пронумерованной так, как её нумерует расширение.
"""

import io
import unittest

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from config import TEST_TABLES_PATH
from MLverification import MLVerification
from PageAnalyzer import PageAnalyzer
from Table import Table
from TableExporter import FORMAT_VERSION, TableExporter, tables_to_xlsx
from TableStructure import StructureTable


def load_tables():
    with open(TEST_TABLES_PATH, "r", encoding="utf-8") as f:
        soup = BeautifulSoup(f, "html.parser")
    return soup.find_all("table", recursive=True)


class test_json(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tables = load_tables()
        cls.first = TableExporter(StructureTable.from_table(Table(cls.tables[0]))).to_dict(
            {'url': 'file:///table_for_test.html', 'index': 0})

    def test_top_level(self):
        self.assertEqual(self.first['format'], FORMAT_VERSION)
        self.assertEqual(self.first['orientation'], 'top')
        self.assertEqual(self.first['size'], {'rows': 15, 'cols': 21})
        self.assertEqual(self.first['source']['index'], 0)

    def test_cell_record(self):
        record = next(cell for cell in self.first['cells']
                      if (cell['row'], cell['col']) == (13, 11))
        self.assertEqual(record['text'], '14200')
        self.assertEqual(record['role'], 'data')
        self.assertEqual(record['type'], 'GENUINE')
        self.assertEqual(record['detail_type'], 'DIMENSION')
        self.assertEqual(record['header_path'][1:], ['🇨🇳 Азия', 'Электроника', 'Ноутбуки'])
        self.assertEqual(record['sidebar_path'], ['📅 2024', '📈 Итого 2024'])
        self.assertNotIn('aggregates', record)

    def test_header_tree(self):
        roots = [node['text'] for node in self.first['header']]
        self.assertEqual(roots, ['Год', 'Квартал', '🌍 Глобальные продажи по регионам',
                                 '📊 Аналитика', '💰 Общий итог'])
        usa = self.first['header'][2]['children'][0]
        self.assertEqual(usa['text'], '🇺🇸 США')
        self.assertEqual([node['text'] for node in usa['children'][0]['children']],
                         ['Смартфоны', 'Ноутбуки'])

    def test_aggregates(self):
        sixth = TableExporter(StructureTable.from_table(Table(self.tables[5]))).to_dict()
        record = next(cell for cell in sixth['cells'] if (cell['row'], cell['col']) == (10, 4))
        self.assertEqual(record['aggregates'], [[7, 4], [8, 4], [9, 4]])

    def test_left_orientation_keeps_page_coordinates(self):
        exported = TableExporter(StructureTable.from_table(Table(self.tables[29]))).to_dict()
        self.assertEqual(exported['orientation'], 'left')
        self.assertEqual(exported['size'], {'rows': 6, 'cols': 2})
        self.assertEqual({(cell['row'], cell['col']) for cell in exported['cells']},
                         {(row, col) for row in range(6) for col in range(2)})


class test_excel(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        tables = load_tables()
        cls.first = StructureTable.from_table(Table(tables[0]))
        cls.thirtieth = StructureTable.from_table(Table(tables[29]))
        cls.exporter = TableExporter(cls.first)
        cls.rotated = TableExporter(cls.thirtieth)

    def _sheet(self, exporter, colored):
        return load_workbook(io.BytesIO(exporter.to_xlsx(colored=colored))).active

    def test_colored(self):
        sheet = self._sheet(self.exporter, colored=True)
        self.assertEqual(sheet.cell(row=14, column=12).value, '14200')
        self.assertEqual(sheet.cell(row=14, column=12).fill.fgColor.rgb, '00000000')
        self.assertEqual(sheet.cell(row=1, column=1).fill.fill_type, 'solid')
        self.assertIn('C1:N1', {str(merged) for merged in sheet.merged_cells.ranges})

    def test_plain(self):
        sheet = self._sheet(self.exporter, colored=False)
        self.assertEqual(sheet.cell(row=14, column=12).value, '14200')
        self.assertIsNone(sheet.cell(row=1, column=1).fill.fill_type)
        self.assertIn('C1:N1', {str(merged) for merged in sheet.merged_cells.ranges})

    def test_rotated_table_is_written_as_on_page(self):
        sheet = self._sheet(self.rotated, colored=False)
        self.assertEqual((sheet.max_row, sheet.max_column), (6, 2))

    def test_several_tables_in_one_file(self):
        """Каждая таблица -- на своём листе с заданным именем"""
        content = tables_to_xlsx([self.first, self.thirtieth],
                                 ['Таблица_1', 'Таблица_30'], colored=False)
        book = load_workbook(io.BytesIO(content))
        self.assertEqual(book.sheetnames, ['Таблица_1', 'Таблица_30'])
        self.assertEqual(book['Таблица_1'].cell(row=14, column=12).value, '14200')
        self.assertEqual(book['Таблица_30'].max_row, 6)


class test_page_analyzer(unittest.TestCase):
    """Ответ сервера по таблицам, пронумерованным как в расширении"""
    @classmethod
    def setUpClass(cls):
        cls.tables = load_tables()
        cls.analyzer = PageAnalyzer(MLVerification())

    def _numbered(self, index):
        tag = self.tables[index]
        for number, cell in enumerate(tag.find_all(['td', 'th'])):
            cell['data-gt-id'] = str(number)
        return str(tag)

    def test_genuine(self):
        answer = self.analyzer.analyze(self._numbered(5), {'index': 5})
        self.assertTrue(answer['genuine'])
        self.assertIsNone(answer['reason'])
        total = next(cell for cell in answer['cells'].values()
                     if cell['lines'][0] == 'Агрегирует ячеек: 13')
        self.assertEqual(total['role'], 'result')
        self.assertEqual(len(total['highlight']), 13)
        self.assertEqual(answer['export']['source'], {'index': 5})

    def test_rejected_by_heuristic(self):
        answer = self.analyzer.analyze(self._numbered(1))
        self.assertFalse(answer['genuine'])
        self.assertIn('LayoutError', answer['reason'])
        self.assertNotIn('cells', answer)

    def test_no_table(self):
        self.assertFalse(self.analyzer.analyze('<p>текст</p>')['genuine'])
